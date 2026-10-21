
import io
import datetime as dt
import urllib.request
import openpyxl

URL = (
    "https://www.cne.gob.mx/da/"
    "Precios_promedio_diarios_y_mensuales_en_estaciones_de_servicio.xlsx"
)

print("DIAGNÓSTICO DE DESCARGA CNE")

req = urllib.request.Request(
    URL,
    headers={
        "User-Agent": "Mozilla/5.0",
        "Cache-Control": "no-cache"
    }
)

with urllib.request.urlopen(req, timeout=90) as response:
    data = response.read()
    print("HTTP:", response.status)
    print("URL final:", response.url)
    print("Tipo de contenido:", response.headers.get("Content-Type"))
    print("Tamaño:", len(data), "bytes")

libro = openpyxl.load_workbook(
    io.BytesIO(data),
    read_only=True,
    data_only=True
)

print("Hojas:", libro.sheetnames)

for hoja in libro:
    fechas = []
    septiembre = 0

    for fila in hoja.iter_rows(values_only=True):
        valor = fila[0] if fila else None

        fecha = None

        if isinstance(valor, dt.datetime):
            fecha = valor.date()
        elif isinstance(valor, dt.date):
            fecha = valor
        elif isinstance(valor, str):
            for formato in ("%d/%m/%Y", "%Y-%m-%d"):
                try:
                    fecha = dt.datetime.strptime(
                        valor.strip(), formato
                    ).date()
                    break
                except ValueError:
                    pass

        if fecha:
            fechas.append(fecha)
            if fecha.year == 2026 and fecha.month == 9:
                septiembre += 1

    print("\nHOJA:", hoja.title)
    print("Última fecha:", max(fechas) if fechas else "No identificada")
    print("Filas septiembre 2026:", septiembre)

libro.close()
