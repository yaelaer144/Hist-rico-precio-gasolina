
"""Descarga e incorpora los precios diarios de gasolina publicados por la CNE."""

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import os
import time
import urllib.request
from pathlib import Path

import openpyxl
from cargar_dof import ROOT

URL = (
    "https://www.cne.gob.mx/da/"
    "Precios_promedio_diarios_y_mensuales_en_estaciones_de_servicio.xlsx"
)

COMBUSTIBLES = ("regular", "premium", "diesel")


def periodo(fecha):
    fin = fecha.replace(day=1) - dt.timedelta(days=1)
    return fin.replace(day=1), fin


def fecha_valida(valor):
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, (int, float)) and 30000 < valor < 90000:
        return openpyxl.utils.datetime.from_excel(valor).date()

    texto = str(valor).strip().split(" ")[0]
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    return None


def descargar():
    ultimo_error = None

    for intento in range(1, 4):
        try:
            solicitud = urllib.request.Request(
                URL,
                headers={
                    "User-Agent": "Mozilla/5.0",
                    "Accept": (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet,*/*"
                    ),
                    "Cache-Control": "no-cache",
                },
            )

            with urllib.request.urlopen(
                solicitud, timeout=90
            ) as respuesta:
                contenido = respuesta.read()
                print(
                    "DIAGNÓSTICO CNE:",
                    "HTTP =", respuesta.status,
                    "URL final =", respuesta.url,
                    "Bytes =", len(contenido),
                    "Tipo =", respuesta.headers.get("Content-Type"),
                    flush=True,
                )

            if not contenido.startswith(b"PK"):
                raise ValueError(
                    "La CNE no devolvió un archivo XLSX válido. "
                    "Puede tratarse de HTML o de un bloqueo."
                )

            return contenido

        except (OSError, ValueError) as error:
            ultimo_error = error
            print(
                f"Intento de descarga {intento}/3: {error}",
                flush=True,
            )
            if intento < 3:
                time.sleep(3 * intento)

    raise ValueError(
        f"No se pudo descargar el Excel de la CNE: {ultimo_error}"
    )


def leer_precios(contenido, inicio, fin):
    libro = openpyxl.load_workbook(
        io.BytesIO(contenido),
        read_only=True,
        data_only=True,
    )

    precios = {}
    ultima_fecha = None

    try:
        print("Hojas encontradas:", libro.sheetnames)

        for hoja in libro:
            columnas = None

            for fila in hoja.iter_rows(values_only=True):
                encabezados = [
                    str(v or "").strip().lower()
                    for v in fila[:12]
                ]

                if (
                    "fecha" in encabezados
                    and any("regular" in x for x in encabezados)
                    and any("premium" in x for x in encabezados)
                    and any(
                        "diésel" in x or "diesel" in x
                        for x in encabezados
                    )
                ):
                    columnas = (
                        encabezados.index("fecha"),
                        next(
                            i for i, x in enumerate(encabezados)
                            if "regular" in x
                        ),
                        next(
                            i for i, x in enumerate(encabezados)
                            if "premium" in x
                        ),
                        next(
                            i for i, x in enumerate(encabezados)
                            if "diésel" in x or "diesel" in x
                        ),
                    )
                    continue

                if columnas is None:
                    continue

                indice_fecha = columnas[0]

                if len(fila) <= indice_fecha:
                    continue

                dia = fecha_valida(fila[indice_fecha])

                if dia is None:
                    continue

                if ultima_fecha is None or dia > ultima_fecha:
                    ultima_fecha = dia

                if not inicio <= dia <= fin:
                    continue

                for combustible, columna in zip(
                    COMBUSTIBLES, columnas[1:]
                ):
                    try:
                        valor = float(
                            str(fila[columna]).replace(",", ".")
                        )

                        if math.isfinite(valor) and 5 <= valor <= 100:
                            precios[
                                (dia.isoformat(), combustible)
                            ] = valor
                    except (ValueError, TypeError, IndexError):
                        pass

    finally:
        libro.close()

    print(
        "DIAGNÓSTICO CNE:",
        "Última fecha del Excel =", ultima_fecha,
        "Precios del periodo =", len(precios),
        "SHA256 =", hashlib.sha256(contenido).hexdigest(),
        flush=True,
    )

    if ultima_fecha is None:
        raise ValueError(
            "No se identificaron fechas diarias en el archivo CNE."
        )

    if ultima_fecha < fin:
        raise ValueError(
            f"El Excel descargado llega hasta {ultima_fecha}, "
            f"pero se necesitan precios hasta {fin}. "
            "Revisar el enlace oficial de descarga."
        )

    esperados = ((fin - inicio).days + 1) * 3

    if len(precios) != esperados:
        faltantes = []

        dia = inicio
        while dia <= fin:
            for combustible in COMBUSTIBLES:
                clave = (dia.isoformat(), combustible)
                if clave not in precios:
                    faltantes.append(clave)
            dia += dt.timedelta(days=1)

        raise ValueError(
            f"Cobertura incompleta: {len(precios)}/{esperados}. "
            f"Primeros registros faltantes: {faltantes[:8]}"
        )

    return precios, ultima_fecha


def incorporar(precios):
    archivo = ROOT / "data/master.json"
    base = json.loads(archivo.read_text(encoding="utf-8"))

    registros = {
        (fila["fecha"], fila["combustible"]): fila
        for fila in base
    }

    for (fecha, combustible), precio in precios.items():
        clave = (fecha, combustible)

        anterior = registros.get(
            clave,
            {"fecha": fecha, "combustible": combustible},
        )

        fila = dict(anterior)
        fila["precio"] = precio
        fila["url_precio"] = URL

        cuota = fila.get("cuota")

        if cuota is not None:
            fila["neto"] = round(precio - cuota, 8)
            fila["fuente"] = "CNE/CRE y DOF"
        else:
            fila["neto"] = None
            fila["fuente"] = "CNE/CRE (precio); cuota pendiente"

        registros[clave] = fila

    resultado = sorted(
        registros.values(),
        key=lambda fila: (fila["fecha"], fila["combustible"]),
    )

    temporal = archivo.with_suffix(".tmp")
    temporal.write_text(
        json.dumps(
            resultado, ensure_ascii=False, allow_nan=False
        ),
        encoding="utf-8",
    )
    os.replace(temporal, archivo)

    with (ROOT / "data/master.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as salida:
        columnas = list(
            dict.fromkeys(
                columna
                for fila in resultado
                for columna in fila
            )
        )
        escritor = csv.DictWriter(salida, fieldnames=columnas)
        escritor.writeheader()
        escritor.writerows(resultado)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--fecha", default=dt.date.today().isoformat()
    )
    parser.add_argument("--dry-run", action="store_true")
    argumentos = parser.parse_args()

    inicio, fin = periodo(
        dt.date.fromisoformat(argumentos.fecha)
    )

    print(f"Consultando CNE del {inicio} al {fin}")

    contenido = descargar()
    precios, ultima_fecha = leer_precios(
        contenido, inicio, fin
    )

    reporte = {
        "inicio": inicio.isoformat(),
        "fin": fin.isoformat(),
        "ultima_fecha_excel": ultima_fecha.isoformat(),
        "precios_validados": len(precios),
        "fuente": URL,
        "sha256": hashlib.sha256(contenido).hexdigest(),
    }

    if argumentos.dry_run:
        print(json.dumps(reporte, ensure_ascii=False, indent=2))
        return

    incorporar(precios)

    (ROOT / "raw").mkdir(parents=True, exist_ok=True)
    (ROOT / "raw/precios.xlsx").write_bytes(contenido)

    (ROOT / "data/actualizacion_mensual.json").write_text(
        json.dumps(reporte, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("Precios CNE incorporados correctamente.")
    print(json.dumps(reporte, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
